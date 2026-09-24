"""
Export & Import Excel untuk data kontrak.
"""
from datetime import date, datetime
from io import BytesIO
import os
from typing import Any, Mapping, Sequence, cast

from flask import (
    Blueprint, flash, redirect, render_template,
    request, session, url_for, send_file
)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.database import execute
from app.middlewares.auth_middleware import login_required


export_bp = Blueprint('export', __name__)


# ============================================================
# KONSTANTA
# ============================================================
MAX_ROWS = 5000  # batas baris per import
MAX_FILE_SIZE = 16 * 1024 * 1024  # 16 MB

ALLOWED_MODES = {'skip', 'replace'}

HEADER_MAP = {
    'NO. KONTRAK': 'no_kontrak',
    'URAIAN PEKERJAAN': 'uraian_pekerjaan',
    'JENIS': 'jenis',
    'LINK DOKUMEN': 'link_dokumen',
    'VENDOR': 'vendor',
    'STATUS PEKERJAAN': 'status_pekerjaan',
    'AWAL KONTRAK': 'awal_kontrak',
    'AKHIR KONTRAK': 'akhir_kontrak',
    'PROGRES BAYAR (%)': 'progres_bayar',
    'PROGRES FISIK (%)': 'progres_fisik',
    'TERMIN': 'termin',
    'UPDATE LKP': 'update_lkp',
    'CATATAN': 'catatan',
    'AKHIR JAMPEL': 'akhir_jampel',
    'AKHIR JAMHAR': 'akhir_jamhar',
    'MASA PEMELIHARAAN': 'masa_pemeliharaan',
    'AKHIR AMANDEMEN': 'akhir_amandemen',
}

# Field yang harus berupa angka
INT_FIELDS = {'progres_bayar', 'progres_fisik'}

# Field yang harus berupa tanggal
DATE_FIELDS = {
    'awal_kontrak', 'akhir_kontrak', 'update_lkp',
    'akhir_jampel', 'akhir_jamhar', 'akhir_amandemen',
}

# Field wajib
REQUIRED_FIELDS = {'no_kontrak', 'awal_kontrak', 'akhir_kontrak'}


# ============================================================
# HELPER
# ============================================================
def _to_date(v):
    """Konversi value apapun ke date atau None."""
    if v is None or v == '':
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d/%m/%Y'):
            try:
                return datetime.strptime(v[:19], fmt).date()
            except ValueError:
                continue
    return None


def _format_date_id(d):
    """Format tanggal Indonesia: 01 Jan 2026."""
    if not d:
        return '-'
    bulan = ['', 'Jan', 'Feb', 'Mar', 'Apr', 'Mei', 'Jun',
             'Jul', 'Agu', 'Sep', 'Okt', 'Nov', 'Des']
    return f"{d.day:02d} {bulan[d.month]} {d.year}"


def _clean_excel_value(val, field=None):
    """Bersihkan nilai dari Excel sesuai field type."""
    if val is None:
        return None

    # Datetime → ISO string
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d')
    if isinstance(val, date):
        return val.strftime('%Y-%m-%d')

    # Angka
    if isinstance(val, (int, float)):
        # Untuk field persen, cast ke int
        if field in INT_FIELDS:
            try:
                return int(round(float(val)))
            except (ValueError, TypeError):
                return 0
        return val

    # String
    if isinstance(val, str):
        val = val.strip()
        if val in ('', '-'):
            return None

        # Kalau field angka, coba parse
        if field in INT_FIELDS:
            try:
                return int(round(float(val.replace('%', '').strip())))
            except (ValueError, TypeError):
                return 0

        # Kalau field tanggal, coba parse
        if field in DATE_FIELDS:
            return _try_parse_date(val)

        return val

    return str(val)


def _try_parse_date(s):
    """Coba parse string sebagai tanggal → ISO format atau None."""
    if not s:
        return None

    s = str(s).strip()

    formats = [
        '%Y-%m-%d',
        '%d/%m/%Y',
        '%d-%m-%Y',
        '%Y/%m/%d',
        '%d %b %Y',
        '%d %B %Y',
        '%d-%b-%Y',
        '%d-%B-%Y',
    ]
    for fmt in formats:
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue

    return None  # bukan tanggal


def _build_export_query():
    return """
        SELECT no_kontrak, uraian_pekerjaan, jenis, link_dokumen, vendor,
               status_pekerjaan, awal_kontrak, akhir_kontrak,
               progres_bayar, progres_fisik, termin, update_lkp, catatan,
               akhir_jampel, akhir_jamhar, masa_pemeliharaan, akhir_amandemen
        FROM kontrak ORDER BY no_kontrak
    """


# ============================================================
# EXPORT EXCEL
# ============================================================
@export_bp.route('/excel')
@login_required
def excel():
    """Download semua kontrak sebagai file Excel."""
    rows = cast(Sequence[Mapping[str, Any]], execute(_build_export_query()))
    today = date.today()

    wb = openpyxl.Workbook()
    ws = cast(Worksheet, wb.active)
    ws.title = "Kontrak"

    headers = [
        'NO. KONTRAK', 'URAIAN PEKERJAAN', 'Jenis', 'Link Dokumen', 'VENDOR',
        'Status Pekerjaan', 'Awal Kontrak', 'Akhir Kontrak',
        'Progres Bayar (%)', 'Progres Fisik (%)', 'Termin', 'Update LKP',
        'Catatan', 'Akhir JamPel', 'Akhir JamHar', 'Masa Pemeliharaan',
        'Akhir Amandemen', 'STATUS Masa Kontrak'
    ]
    ws.append(headers)

    # Format header
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill('solid', fgColor='34495E')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.row_dimensions[1].height = 30

    # Isi data
    for r in rows:
        akhir_kontrak = _to_date(r['akhir_kontrak'])
        akhir_amandemen = _to_date(r['akhir_amandemen'])
        akhir_efektif = akhir_amandemen or akhir_kontrak

        status_masa = '-'
        if akhir_efektif:
            sisa = (akhir_efektif - today).days
            if sisa < 0:
                status_masa = f'Berakhir ({-sisa} hari lalu)'
            elif sisa <= 30:
                status_masa = f'Akan Berakhir ({sisa} hari)'
            else:
                status_masa = f'Aktif ({sisa} hari)'

        ws.append([
            r['no_kontrak'] or '',
            r['uraian_pekerjaan'] or '',
            r['jenis'] or '',
            r['link_dokumen'] or '',
            r['vendor'] or '',
            r['status_pekerjaan'] or '',
            _format_date_id(_to_date(r['awal_kontrak'])),
            _format_date_id(akhir_kontrak),
            int(r['progres_bayar'] or 0),
            int(r['progres_fisik'] or 0),
            r['termin'] or '',
            _format_date_id(_to_date(r['update_lkp'])),
            r['catatan'] or '',
            _format_date_id(_to_date(r['akhir_jampel'])),
            _format_date_id(_to_date(r['akhir_jamhar'])),
            r['masa_pemeliharaan'] or '',
            _format_date_id(akhir_amandemen),
            status_masa,
        ])

    # Auto width
    for i, col in enumerate(ws.columns, 1):
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 40)

    # Freeze header + autofilter
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f'kontrak_{today.strftime("%Y%m%d_%H%M%S")}.xlsx'
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


# ============================================================
# TEMPLATE EXCEL
# ============================================================
@export_bp.route('/template')
@login_required
def template_excel():
    """Download template Excel kosong untuk import."""
    wb = openpyxl.Workbook()
    ws = cast(Worksheet, wb.active)
    ws.title = "Template"

    headers = [
        'NO. KONTRAK', 'URAIAN PEKERJAAN', 'Jenis', 'Link Dokumen', 'VENDOR',
        'Status Pekerjaan', 'Awal Kontrak', 'Akhir Kontrak',
        'Progres Bayar (%)', 'Progres Fisik (%)', 'Termin', 'Update LKP',
        'Catatan', 'Akhir JamPel', 'Akhir JamHar', 'Masa Pemeliharaan',
        'Akhir Amandemen'
    ]
    ws.append(headers)

    # Format header
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='34495E')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Contoh baris
    ws.append([
        'KTR-2026-001', 'Pengadaan Alat Tulis Kantor', 'Jasa',
        'https://drive.google.com/...', 'PT Maju Jaya',
        'Berjalan', '01/01/2026', '01/12/2026',
        50, 60, 'Termin 1', '15/06/2026',
        'Catatan contoh', '01/12/2026', '01/12/2026', '90 hari', ''
    ])

    # Auto width
    for i, col in enumerate(ws.columns, 1):
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 30)

    ws.freeze_panes = 'A2'

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        as_attachment=True,
        download_name='template_import_kontrak.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


# ============================================================
# IMPORT EXCEL
# ============================================================
@export_bp.route('/import', methods=['GET'])
@login_required
def import_page():
    """Halaman form import Excel."""
    return render_template('kontrak/import.html')


@export_bp.route('/import/excel', methods=['POST'])
@login_required
def import_excel():
    """Proses upload & import file Excel."""
    # ---------- Validasi file ----------
    if 'file' not in request.files:
        flash('Tidak ada file yang diupload', 'danger')
        return redirect(url_for('export.import_page'))

    file = request.files['file']
    if not file or file.filename == '':
        flash('File belum dipilih', 'danger')
        return redirect(url_for('export.import_page'))

    # Cek ekstensi
    allowed = {'.xlsx', '.xls'}
    filename = file.filename or ''
    ext = os.path.splitext(filename)[1].lower()
    if ext not in allowed:
        flash(f'Format file harus {", ".join(allowed)}', 'danger')
        return redirect(url_for('export.import_page'))

    # Cek ukuran file
    file.seek(0, os.SEEK_END)
    size = file.tell()
    file.seek(0)
    if size > MAX_FILE_SIZE:
        flash(f'File terlalu besar (max {MAX_FILE_SIZE // 1024 // 1024} MB)', 'danger')
        return redirect(url_for('export.import_page'))

    # Mode
    mode = request.form.get('mode', 'skip')
    if mode not in ALLOWED_MODES:
        mode = 'skip'

    # ---------- Baca file ----------
    try:
        wb = openpyxl.load_workbook(BytesIO(file.read()), data_only=True)
        ws = cast(Worksheet, wb.active)
        if ws is None:
            flash('File Excel tidak memiliki worksheet aktif', 'danger')
            return redirect(url_for('export.import_page'))

        # Header
        header_row = [str(c.value or '').strip().upper() for c in ws[1]]

        col_index = {}
        for i, h in enumerate(header_row):
            if h in HEADER_MAP:
                col_index[HEADER_MAP[h]] = i

        if 'no_kontrak' not in col_index:
            flash('Kolom "NO. KONTRAK" tidak ditemukan di file Excel', 'danger')
            return redirect(url_for('export.import_page'))

        # ---------- Proses baris ----------
        imported = 0
        updated = 0
        skipped = 0
        errors = []
        seen_in_file = set()  # deteksi duplikat dalam 1 file

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            # Batas baris
            if row_idx > MAX_ROWS + 1:
                errors.append(f'Berhenti di baris {MAX_ROWS}: max {MAX_ROWS} baris per import')
                break

            # Skip baris kosong
            if not row or not any(row):
                continue

            # Ambil data
            data = {}
            for field, idx in col_index.items():
                if idx < len(row):
                    data[field] = _clean_excel_value(row[idx], field)

            # Validasi no_kontrak
            no_kontrak = (data.get('no_kontrak') or '').strip()
            if not no_kontrak:
                skipped += 1
                continue

            # Deteksi duplikat dalam file
            if no_kontrak in seen_in_file:
                errors.append(f'Baris {row_idx}: duplikat "{no_kontrak}" dalam file')
                skipped += 1
                continue
            seen_in_file.add(no_kontrak)

            # Validasi field wajib
            if not data.get('awal_kontrak'):
                errors.append(f'Baris {row_idx}: Awal Kontrak kosong / format salah')
                skipped += 1
                continue
            if not data.get('akhir_kontrak'):
                errors.append(f'Baris {row_idx}: Akhir Kontrak kosong / format salah')
                skipped += 1
                continue

            # Default value
            data.setdefault('progres_bayar', 0)
            data.setdefault('progres_fisik', 0)
            data.setdefault('status_pekerjaan', 'Berjalan')

            # ---------- Cek duplikat di DB ----------
            existing = execute(
                "SELECT id FROM kontrak WHERE no_kontrak=?", (no_kontrak,)
            )

            if existing:
                if mode == 'skip':
                    skipped += 1
                    continue
                elif mode == 'replace':
                    update_data = {k: v for k, v in data.items() if k != 'no_kontrak'}
                    if update_data:
                        try:
                            sets = ', '.join([f"{k}=?" for k in update_data.keys()])
                            execute(
                                f"UPDATE kontrak SET {sets} WHERE no_kontrak=?",
                                list(update_data.values()) + [no_kontrak],
                                fetch=False
                            )
                            updated += 1
                        except Exception as e:
                            errors.append(f'Baris {row_idx}: {str(e)[:80]}')
                            skipped += 1
                    else:
                        skipped += 1
                    continue

            # ---------- Insert baru ----------
            try:
                cols = ', '.join(data.keys())
                placeholders = ', '.join(['?'] * len(data))
                execute(
                    f"INSERT INTO kontrak ({cols}) VALUES ({placeholders})",
                    list(data.values()),
                    fetch=False
                )
                imported += 1
            except Exception as e:
                errors.append(f'Baris {row_idx}: {str(e)[:80]}')
                skipped += 1

        # ---------- Log ----------
        try:
            from app.repositories.log_repo import LogRepository
            LogRepository.add(
                session.get('user_id'),
                'IMPORT',
                f"Import Excel: {imported} baru, {updated} diperbarui, {skipped} dilewati"
            )
        except Exception:
            pass

        # ---------- Sync ke Sheets ----------
        try:
            from app.services.sheets_service import SheetsService
            SheetsService.sync_all()
        except Exception:
            pass

        # ---------- Flash ----------
        msg = f'Import selesai: {imported} baru, {updated} diperbarui, {skipped} dilewati'
        flash(msg, 'success')

        for err in errors[:5]:
            flash(err, 'warning')

        if len(errors) > 5:
            flash(f'... dan {len(errors) - 5} error lainnya', 'warning')

        return redirect(url_for('kontrak.list'))

    except Exception as e:
        flash(f'Gagal import: {str(e)[:200]}', 'danger')
        return redirect(url_for('export.import_page'))